"""Build the fixed, synthetic Study acceptance corpus on Windows."""
from __future__ import annotations

import hashlib
import io
import json
import sys
import tempfile
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from prepare import convert_pptx

HERE = Path(__file__).resolve().parent
FIXTURES = HERE / 'fixtures'


def text_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page(width=640, height=420)
    cjk = fitz.Font(fontname='cjk')
    page.insert_font(fontname='StudyCJK', fontbuffer=cjk.buffer)
    page.insert_text((42, 55), 'Chapter 1: Energy  能量', fontname='StudyCJK', fontsize=22)
    page.insert_text((42, 110), 'RedAnswer', fontsize=18, color=(0.95, 0.1, 0.1))
    page.insert_text((42, 155), 'BoldAnswer', fontname='hebo', fontsize=18)
    page.insert_text((42, 200), 'UnderlinedAnswer', fontsize=18)
    line = page.search_for('UnderlinedAnswer')[0]
    page.draw_line((line.x0, line.y1 + 1), (line.x1, line.y1 + 1), color=(0, 0, 0), width=1.2)
    page.insert_text((42, 245), 'HighlightedAnswer', fontsize=18)
    page.add_highlight_annot(page.search_for('HighlightedAnswer'))
    page.insert_text((42, 300), 'Normal context must stay visible.', fontsize=16)
    page = doc.new_page(width=420, height=640)
    page.insert_font(fontname='StudyCJK', fontbuffer=cjk.buffer)
    page.insert_text((40, 62), 'Portrait page  纵向页面', fontname='StudyCJK', fontsize=20)
    page.insert_text((40, 115), 'RepeatedWord RepeatedWord', fontsize=16)
    doc.set_metadata({'title': 'Study synthetic text corpus', 'creator': 'study-app'})
    doc.subset_fonts()
    doc.save(path, garbage=4, deflate=True)
    doc.close()


def scan_pdf(path: Path) -> dict:
    vector = fitz.open()
    page = vector.new_page(width=640, height=420)
    page.insert_text((45, 65), 'Alpha', fontsize=26, color=(0.95, 0.1, 0.1))
    page.draw_rect((42, 96, 180, 139), color=None, fill=(1, 0.9, 0.28))
    page.insert_text((45, 128), 'Beta', fontsize=26)
    page.insert_text((45, 193), 'Gamma', fontsize=26)
    page.draw_line((45, 200), (150, 200), color=(0.06, 0.06, 0.06), width=2)
    page.insert_text((45, 255), 'Delta', fontsize=26)
    boxes = {}
    for word in ('Alpha', 'Beta', 'Gamma', 'Delta'):
        box = page.search_for(word)[0]
        boxes[word] = [round(box.x0 / 640, 4), round(box.y0 / 420, 4),
                       round(box.width / 640, 4), round(box.height / 420, 4)]
    image = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes('png')
    vector.close()
    scan = fitz.open()
    page = scan.new_page(width=640, height=420)
    page.insert_image(page.rect, stream=image)
    scan.set_metadata({'title': 'Study synthetic raster scan', 'creator': 'study-app'})
    scan.save(path, garbage=4, deflate=True)
    scan.close()
    return boxes


def large_pdf(path: Path) -> None:
    doc = fitz.open()
    for index in range(120):
        page = doc.new_page(width=595, height=842)
        page.insert_text((40, 55), f'Study large document - page {index + 1:03d}', fontsize=18)
        for row in range(18):
            page.insert_text((42, 100 + row * 31), f'Chapter {index // 12 + 1} / paragraph {row + 1}: keyword and answer.', fontsize=11)
    doc.set_metadata({'title': 'Study synthetic 120 page document', 'creator': 'study-app'})
    doc.save(path, garbage=4, deflate=True)
    doc.close()


def complex_pptx(path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    title = slide.shapes.add_textbox(Inches(.65), Inches(.45), Inches(11), Inches(.8))
    run = title.text_frame.paragraphs[0].add_run()
    run.text = '第一章 Energy / 能量'; run.font.name = 'Microsoft YaHei'; run.font.size = Pt(30)
    box = slide.shapes.add_textbox(Inches(.8), Inches(1.8), Inches(6), Inches(1.2))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = 'RedAnswer'; run.font.size = Pt(27); run.font.color.rgb = RGBColor(220, 38, 38)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = '  UnderlinedAnswer'; run.font.size = Pt(27); run.font.underline = True
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.8), Inches(3.6), Inches(3.2), Inches(1.2))

    slide = prs.slides.add_slide(blank)
    table = slide.shapes.add_table(3, 3, Inches(.8), Inches(1.1), Inches(6), Inches(2.6)).table
    for row, values in enumerate((('Item', 'A', 'B'), ('Speed', '12', '15'), ('Power', '7', '9'))):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    art = fitz.open(); art_page = art.new_page(width=360, height=200)
    art_page.draw_circle((100, 100), 70, color=(0.1, .35, .8), fill=(.7, .85, 1))
    art_page.insert_text((195, 105), 'E=mc2', fontsize=30)
    png = art_page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes('png'); art.close()
    slide.shapes.add_picture(io.BytesIO(png), Inches(7.1), Inches(1.2), width=Inches(5))

    slide = prs.slides.add_slide(blank)
    chart = CategoryChartData(); chart.categories = ['一月', '二月', '三月']
    chart.add_series('甲', (10, 18, 12)); chart.add_series('乙', (14, 11, 21))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(.8),
                           Inches(11), Inches(5.8), chart)

    slide = prs.slides.add_slide(blank)
    slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.7), Inches(1), Inches(8), Inches(4))
    box = slide.shapes.add_textbox(Inches(1.1), Inches(2), Inches(7), Inches(1.4))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = 'BoldAnswer  重要结论'; run.font.name = 'Microsoft YaHei'; run.font.bold = True; run.font.size = Pt(26)
    prs.save(path)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build() -> dict:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    paths = {name: FIXTURES / name for name in (
        'text_styles.pdf', 'scan_styles.pdf', 'large_120_pages.pdf',
        'deck_complex.pptx', 'deck_complex_powerpoint.pdf')}
    text_pdf(paths['text_styles.pdf'])
    scan_boxes = scan_pdf(paths['scan_styles.pdf'])
    large_pdf(paths['large_120_pages.pdf'])
    complex_pptx(paths['deck_complex.pptx'])
    with tempfile.TemporaryDirectory(prefix='study-corpus-') as temp:
        convert_pptx(paths['deck_complex.pptx'], paths['deck_complex_powerpoint.pdf'], Path(temp))
    manifest = {
        'schema': 'study-quality-corpus/1', 'synthetic': True,
        'reference': 'Windows PowerPoint 16 PDF export',
        'fixtures': {
            name: {'sha256': sha256(path), 'pages': 4 if name.endswith('.pptx') or name.startswith('deck_') else
                   120 if name.startswith('large_') else 2 if name.startswith('text_') else 1}
            for name, path in paths.items()
        },
        'scan_expected': {
            'keyword': ['Alpha', 'Beta', 'Gamma', 'Delta'],
            'red_text': ['Alpha'], 'highlight': ['Beta'], 'underline': ['Gamma'],
            'negative': ['Delta'],
            'boxes': scan_boxes,
        },
        'pptx_provisional_gates': {
            'page_count_equal': True, 'max_aspect_ratio_delta': 0.005,
            'mean_rgb_mae_max': 12.0, 'worst_page_rgb_mae_max': 20.0,
            'min_text_character_recall': 0.98,
        },
    }
    (HERE / 'corpus.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return manifest


if __name__ == '__main__':
    print(json.dumps(build(), ensure_ascii=False, indent=2))
