import hashlib
import json
import tempfile
import unittest
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from prepare import prepare, pptx_marked_text, normalize_pdf

class PrepareTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
    def tearDown(self): self.tmp.cleanup()
    def fixture(self):
        path=self.root/'source.pdf'
        doc=fitz.open();p=doc.new_page(width=600,height=400)
        p.insert_text((40,60),'Normal textbook content',fontsize=18)
        p.insert_text((40,110),'ColorAnswer',fontsize=18,color=(0,0,1))
        p.insert_text((40,160),'BoldAnswer',fontsize=18,fontname='hebo')
        p.insert_text((40,210),'KeywordAnswer',fontsize=18)
        doc.save(path);doc.close()
        return path
    def test_rules_union_and_export_removes_answers(self):
        source=self.fixture();original=source.read_bytes()
        m=prepare(source,self.root/'out',keywords=['KeywordAnswer'],colors=[0x0000FF],bold=True)
        self.assertEqual(len(m['pages'][0]['masks']),3)
        self.assertEqual(source.read_bytes(),original)
        pdf=self.root/'out/textbook.pdf'
        self.assertEqual(hashlib.sha256(pdf.read_bytes()).hexdigest(),m['pdf_sha256'])
        with fitz.open(self.root/'out/practice.pdf') as d:
            text=d[0].get_text()
            self.assertIn('Normal textbook content',text)
            for answer in ['ColorAnswer','BoldAnswer','KeywordAnswer']:self.assertNotIn(answer,text)
        for mask in m['pages'][0]['masks']:
            x,y,w,h=mask['rect'];self.assertGreater(w,0);self.assertGreater(h,0)
            self.assertTrue(0<=x<x+w<=1);self.assertTrue(0<=y<y+h<=1)
    def test_overlapping_rules_deduplicate(self):
        source=self.fixture()
        m=prepare(source,self.root/'out',keywords=['ColorAnswer','ColorAnswer'],colors=[0x0000FF])
        self.assertEqual(len(m['pages'][0]['masks']),1)
    def test_no_match_is_valid(self):
        m=prepare(self.fixture(),self.root/'out',keywords=['absent'])
        self.assertEqual(m['pages'][0]['masks'],[])
    def test_dont_overwrite_previous_output(self):
        source=self.fixture();prepare(source,self.root/'out',bold=True)
        with self.assertRaises(ValueError):prepare(source,self.root/'out',bold=True)
    def test_scanned_pdf_reports_no_text(self):
        source=self.root/'scan.pdf';doc=fitz.open();p=doc.new_page()
        p.draw_rect(fitz.Rect(40,40,100,100),fill=(0.2,0.2,0.2))
        doc.save(source);doc.close()
        m=prepare(source,self.root/'out',keywords=['answer'])
        self.assertEqual(len(m['warnings']),1)
        self.assertEqual(m['pages'][0]['masks'],[])
    def test_blank_page_is_preserved(self):
        source=self.root/'blank.pdf';doc=fitz.open();doc.new_page();doc.save(source);doc.close()
        m=prepare(source,self.root/'out',keywords=['answer'])
        self.assertEqual(len(m['pages']),1)
        self.assertEqual(m['pages'][0]['masks'],[])
    def test_pdf_underline_is_not_silently_guessed(self):
        with self.assertRaises(ValueError):prepare(self.fixture(),self.root/'out',underline=True)
    def test_pptx_conversion_and_explicit_underline(self):
        source=self.root/'教材.pptx';prs=Presentation()
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(1))
        run=box.text_frame.paragraphs[0].add_run();run.text='UnderlinedAnswer'
        run.font.size=Pt(24);run.font.underline=True;run.font.color.rgb=RGBColor(255,0,0)
        prs.save(source)
        marked,_=pptx_marked_text(source,underline=True)
        self.assertEqual(marked,[['UnderlinedAnswer']])
        m=prepare(source,self.root/'输出',underline=True)
        self.assertEqual(len(m['pages']),1)
        self.assertGreaterEqual(len(m['pages'][0]['masks']),1)
        self.assertTrue((self.root/'输出/practice.pdf').exists())

    def test_rotated_and_cropped_pdf_keeps_visible_geometry(self):
        for rotation in [0, 90, 180, 270]:
            for cropped in [False, True]:
                with self.subTest(rotation=rotation, cropped=cropped):
                    source=self.root/f'geometry-{rotation}-{cropped}.pdf'
                    doc=fitz.open(); page=doc.new_page(width=600,height=400)
                    page.insert_text((100,120),'GeometryAnswer',fontsize=18)
                    if cropped: page.set_cropbox(fitz.Rect(40,40,560,360))
                    page.set_rotation(rotation)
                    expected=page.search_for('GeometryAnswer')[0]*page.rotation_matrix
                    size=page.rect
                    doc.save(source);doc.close()
                    out=self.root/f'geometry-out-{rotation}-{cropped}'
                    manifest=prepare(source,out,keywords=['GeometryAnswer'])
                    with fitz.open(out/'textbook.pdf') as normalized:
                        page=normalized[0]
                        self.assertEqual(page.rotation,0)
                        self.assertEqual(page.rect,size)
                        actual=page.search_for('GeometryAnswer')[0]
                        for a,b in zip(expected,actual): self.assertAlmostEqual(a,b,places=3)
                        self.assertEqual(len(manifest['pages'][0]['masks']),1)
    def test_annotation_appearance_is_not_dropped(self):
        source=self.fixture()
        with fitz.open(source) as doc:
            page=doc[0];page.add_highlight_annot(page.search_for('ColorAnswer'))
            doc.save(self.root/'annotated.pdf')
        normalize_pdf(self.root/'annotated.pdf',self.root/'normalized.pdf')
        with fitz.open(self.root/'annotated.pdf') as before, fitz.open(self.root/'normalized.pdf') as after:
            original=before[0].get_pixmap().samples
            copied=after[0].get_pixmap().samples
            mean_error=sum(abs(a-b) for a,b in zip(original,copied))/len(original)
            self.assertLess(mean_error,0.1)
            self.assertIn('ColorAnswer',after[0].get_text())

if __name__=='__main__':unittest.main(verbosity=2)
